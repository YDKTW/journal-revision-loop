#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把審查與修訂流程輸出的 final_manuscript，依「版面規格」轉成可投稿用的 Word 檔（.docx），
並可選擇另存 .doc、.pdf（需要本機安裝 LibreOffice）。

版面規格（layout_spec）來源：
  1. 流程的輸出變數 layout_spec（使用者上傳期刊稿約時，由格式解析節點整理；稿約沒寫明的項目一律用內建預設）
  2. 或用 --spec 指定一個 JSON 檔
  3. 都沒有時，使用內建預設（標楷體／Times New Roman、12 pt、1.5 倍行高、A4）

用法：
  python export_docx.py result.json -o 修訂稿.docx            # result.json＝Dify 輸出（含 final_manuscript、layout_spec）
  python export_docx.py 稿件.txt --spec spec.json -o 修訂稿.docx
  python export_docx.py result.json -o 修訂稿.docx --also-doc --also-pdf
  python export_docx.py result.json -o 投稿版.docx --strip-markers   # 移除【待作者…】標記（定稿用，慎用）

預設會保留【待作者補充／確認】標記並以黃底標示，避免未處理的缺口被悄悄帶進投稿稿。
Google 文件：把 .docx 上傳到 Google 雲端硬碟，右鍵「開啟方式 → Google 文件」即可。
"""
import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

DEFAULT_LAYOUT = {
    "page": {"size": "A4", "margin_top_cm": 2.54, "margin_bottom_cm": 2.54, "margin_left_cm": 3.17, "margin_right_cm": 3.17},
    "font_zh": "標楷體",
    "font_en": "Times New Roman",
    "body": {"size_pt": 12, "line_spacing": 1.5, "first_line_indent_chars": 2, "align": "justify"},
    "title": {"size_pt": 16, "bold": True, "align": "center"},
    "heading1": {"size_pt": 14, "bold": True, "align": "left"},
    "heading2": {"size_pt": 13, "bold": True, "align": "left"},
    "heading3": {"size_pt": 12, "bold": True, "align": "left"},
    "table_caption": {"position": "above", "size_pt": 11, "bold": True, "align": "center"},
    "table": {"size_pt": 10, "style": "grid"},
    "references": {"size_pt": 12, "line_spacing": 1.0, "hanging_indent_chars": 2},
    "page_number": {"enabled": True, "position": "center"},
}

TODO_RE = re.compile(r"(〔?【(?:待作者補充|待作者確認|建議補充)[^】]*】〕?)")
LOCK_TAG_RE = re.compile(r"^\s*<!--/?LOCK-->\s*$")
CJK_NUM = "一二三四五六七八九十"
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


# ------------------------------------------------------------------ 版面規格

def merge_layout(base, spec):
    """spec 中非 null 的值覆蓋 base；回傳新字典。"""
    out = copy.deepcopy(base)
    if not isinstance(spec, dict):
        return out
    for k, v in spec.items():
        if k in ("unspecified", "notes"):
            out[k] = v
        elif isinstance(v, dict) and isinstance(out.get(k), dict):
            for kk, vv in v.items():
                if vv is not None:
                    out[k][kk] = vv
        elif v is not None:
            out[k] = v
    return out


def load_layout(spec):
    if isinstance(spec, str):
        spec = spec.strip()
        if not spec:
            return copy.deepcopy(DEFAULT_LAYOUT)
        try:
            spec = json.loads(spec)
        except Exception:
            return copy.deepcopy(DEFAULT_LAYOUT)
    return merge_layout(DEFAULT_LAYOUT, spec)


# ------------------------------------------------------------------ 文字與段落格式

def _fonts(run, zh, en, size, bold=None, italic=None):
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), en)
    rf.set(qn("w:eastAsia"), zh)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _para(doc, align="left", line_spacing=1.0, indent_chars=0, size=12, before=0, after=0,
          keep_next=False, hanging_chars=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    if line_spacing:
        pf.line_spacing = float(line_spacing)
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if indent_chars:
        pf.first_line_indent = Pt(size * float(indent_chars))
    if hanging_chars:
        pf.left_indent = Pt(size * float(hanging_chars))
        pf.first_line_indent = Pt(-size * float(hanging_chars))
    pf.keep_with_next = keep_next
    return p


def _runs(p, text, L, size, bold=None, keep_markers=True, bold_prefix=None):
    """依【待作者…】標記切段；標記以黃底標示（或移除）。"""
    zh, en = L["font_zh"], L["font_en"]
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        _fonts(r, zh, en, size, bold=True)
        text = text[len(bold_prefix):]
    for part in TODO_RE.split(text):
        if not part:
            continue
        if TODO_RE.fullmatch(part):
            if not keep_markers:
                continue
            r = p.add_run(part)
            _fonts(r, zh, en, size, bold=bold)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = p.add_run(part)
            _fonts(r, zh, en, size, bold=bold)


def _page_number(section, L):
    cfg = L.get("page_number") or {}
    if not cfg.get("enabled"):
        return
    p = section.footer.paragraphs[0]
    pos = str(cfg.get("position") or "center")
    p.alignment = ALIGN.get({"left": "left", "right": "right"}.get(pos, "center"))
    r = p.add_run()
    _fonts(r, L["font_zh"], L["font_en"], 10)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tag)
            r._element.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._element.append(it)


# ------------------------------------------------------------------ 表格

def _borders(cell_or_table_pr, **edges):
    """edges: top/bottom/left/right = (size, val) 或 None"""
    b = OxmlElement("w:tblBorders" if cell_or_table_pr.tag.endswith("tblPr") else "w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        spec = edges.get(edge)
        el = OxmlElement("w:" + edge)
        if spec:
            el.set(qn("w:val"), spec[1])
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    cell_or_table_pr.append(b)


def _add_table(doc, rows, L, keep_markers):
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tcfg = L.get("table") or {}
    size = tcfg.get("size_pt", 10)
    style = str(tcfg.get("style") or "grid")
    if style == "three_line":
        _borders(t._tbl.tblPr, top=(12, "single"), bottom=(12, "single"))
    else:
        t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = t.cell(i, j)
            txt = row[j] if j < len(row) else ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            _runs(p, txt, L, size, bold=(i == 0), keep_markers=keep_markers)
            if style == "three_line" and i == 0:
                _borders(cell._tc.get_or_add_tcPr(), bottom=(6, "single"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# ------------------------------------------------------------------ 稿件解析

def _heading_level(line):
    t = line.strip()
    if not t or len(t) > 45 or t.endswith(("。", "；", "：")):
        return 0
    m = re.match(r"^(#{1,3})\s+(.*)$", t)
    if m:
        return len(m.group(1))
    if t in ("摘要", "Abstract", "參考文獻", "References", "附錄") or re.match(rf"^[{CJK_NUM}]+、", t):
        return 1
    if re.match(rf"^（[{CJK_NUM}]+）", t):
        return 2
    if re.match(r"^（\d+）|^\d+\.\d+", t):
        return 3
    return 0


def parse_manuscript(text):
    """回傳區塊清單：('title'|'h1'|'h2'|'h3'|'caption'|'table'|'keywords'|'ref'|'p', 內容)。"""
    lines = [ln.rstrip() for ln in (text or "").replace("\r", "").split("\n")]
    lines = [ln for ln in lines if not LOCK_TAG_RE.match(ln)]
    blocks, i, in_refs, seen_title = [], 0, False, False
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        # 表格（Markdown 的 | 或 Tab 分隔）
        if s.startswith("|") or "\t" in ln:
            rows = []
            while i < len(lines) and (lines[i].strip().startswith("|") or "\t" in lines[i]):
                cur = lines[i].strip()
                if not re.fullmatch(r"\|?[\s:|-]+\|?", cur):
                    rows.append([c.strip() for c in (cur.strip("|").split("|") if cur.startswith("|") else lines[i].split("\t"))])
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue
        i += 1
        s_clean = re.sub(r"^#{1,6}\s+", "", s)
        if not seen_title:
            blocks.append(("title", s_clean))
            seen_title = True
            continue
        lvl = _heading_level(s)
        if lvl:
            in_refs = s_clean in ("參考文獻", "References")
            blocks.append(("h%d" % lvl, s_clean))
            continue
        if re.match(r"^(表|圖)\s*\d+", s_clean) and len(s_clean) <= 60:
            blocks.append(("caption", s_clean))
            continue
        if re.match(r"^(關鍵詞|關鍵字|Keywords?)\s*[:：]", s_clean):
            blocks.append(("keywords", s_clean))
            continue
        blocks.append(("ref" if in_refs else "p", s_clean))
    # 圖表標題與表格配對（標題原位於表格之前；依規格決定上或下）
    return blocks


# ------------------------------------------------------------------ 主流程

def build_docx(manuscript_text, layout, out_path, keep_markers=True):
    L = load_layout(layout)
    blocks = parse_manuscript(manuscript_text)
    doc = Document()

    # 頁面
    sec = doc.sections[0]
    page = L["page"]
    if str(page.get("size", "A4")).upper() == "LETTER":
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(page["margin_top_cm"]), Cm(page["margin_bottom_cm"])
    sec.left_margin, sec.right_margin = Cm(page["margin_left_cm"]), Cm(page["margin_right_cm"])
    _page_number(sec, L)

    # 預設樣式
    st = doc.styles["Normal"]
    st.font.name = L["font_en"]
    st.element.rPr.rFonts.set(qn("w:eastAsia"), L["font_zh"])
    st.font.size = Pt(L["body"]["size_pt"])

    body, refs, cap = L["body"], L["references"], L["table_caption"]
    title_text = ""
    # 處理圖表標題位置：把「標題＋緊接的表格」成對輸出
    k = 0
    while k < len(blocks):
        kind, val = blocks[k]
        if kind == "title":
            title_text = val
            c = L["title"]
            p = _para(doc, c.get("align", "center"), 1.2, 0, c["size_pt"], 0, 12, True)
            _runs(p, val, L, c["size_pt"], bold=c.get("bold", True), keep_markers=keep_markers)
        elif kind in ("h1", "h2", "h3"):
            c = L["heading" + kind[1]]
            before = {"h1": 14, "h2": 10, "h3": 6}[kind]
            p = _para(doc, c.get("align", "left"), 1.2, 0, c["size_pt"], before, 6, True)
            _runs(p, val, L, c["size_pt"], bold=c.get("bold", True), keep_markers=keep_markers)
        elif kind == "caption":
            nxt = blocks[k + 1] if k + 1 < len(blocks) else None
            if nxt and nxt[0] == "table":
                pos = str(cap.get("position") or "above")
                def put_caption():
                    p = _para(doc, cap.get("align", "center"), 1.0, 0, cap["size_pt"], 6, 4, pos == "above")
                    _runs(p, val, L, cap["size_pt"], bold=cap.get("bold", True), keep_markers=keep_markers)
                if pos == "above":
                    put_caption()
                    _add_table(doc, nxt[1], L, keep_markers)
                else:
                    _add_table(doc, nxt[1], L, keep_markers)
                    put_caption()
                k += 1
            else:
                p = _para(doc, cap.get("align", "center"), 1.0, 0, cap["size_pt"], 6, 4)
                _runs(p, val, L, cap["size_pt"], bold=cap.get("bold", True), keep_markers=keep_markers)
        elif kind == "table":
            _add_table(doc, val, L, keep_markers)
        elif kind == "keywords":
            p = _para(doc, "left", body["line_spacing"], 0, body["size_pt"], 6, 6)
            _runs(p, val, L, body["size_pt"], keep_markers=keep_markers, bold_prefix=re.match(r"^[^:：]+[:：]", val).group(0))
        elif kind == "ref":
            p = _para(doc, "left", refs.get("line_spacing", 1.0), 0, refs["size_pt"], 0, 4, False, refs.get("hanging_indent_chars", 2))
            _runs(p, val, L, refs["size_pt"], keep_markers=keep_markers)
        else:
            p = _para(doc, body.get("align", "justify"), body["line_spacing"], body.get("first_line_indent_chars", 2),
                      body["size_pt"], 0, 0)
            _runs(p, val, L, body["size_pt"], keep_markers=keep_markers)
        k += 1

    if title_text:
        doc.core_properties.title = title_text
    doc.save(out_path)
    return out_path


def _soffice_convert(src, fmt):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None, "找不到 LibreOffice（soffice），無法轉成 .%s；請改在 Word 中「另存新檔」。" % fmt
    outdir = str(Path(src).parent)
    try:
        subprocess.run([exe, "--headless", "--convert-to", fmt, "--outdir", outdir, str(src)],
                       check=True, capture_output=True, timeout=180)
    except Exception as e:  # noqa: BLE001
        return None, "轉成 .%s 失敗：%s" % (fmt, e)
    out = Path(src).with_suffix("." + fmt)
    return (str(out), "") if out.exists() else (None, "轉成 .%s 失敗：沒有產生檔案" % fmt)


def pick_outputs(data):
    """從 Dify 的輸出 JSON 中取出 final_manuscript 與 layout_spec（支援多種外層結構）。"""
    if isinstance(data, dict):
        for path in (("outputs",), ("data", "outputs"), ()):
            d = data
            for key in path:
                d = d.get(key, {}) if isinstance(d, dict) else {}
            if isinstance(d, dict) and d.get("final_manuscript"):
                return d.get("final_manuscript", ""), d.get("layout_spec", "")
    return "", ""


def main():
    ap = argparse.ArgumentParser(description="把修訂後的稿件依期刊版面規格輸出為 Word 檔")
    ap.add_argument("input", help="Dify 輸出 JSON（含 final_manuscript），或純文字稿件（.txt／.md）")
    ap.add_argument("-o", "--out", default="修訂稿_依期刊格式.docx")
    ap.add_argument("--spec", help="版面規格 JSON 檔（未提供則使用輸出 JSON 內的 layout_spec，再不然用內建預設）")
    ap.add_argument("--strip-markers", action="store_true", help="移除【待作者補充／確認】標記（定稿用，慎用）")
    ap.add_argument("--also-doc", action="store_true", help="另存 .doc（需 LibreOffice）")
    ap.add_argument("--also-pdf", action="store_true", help="另存 .pdf（需 LibreOffice）")
    a = ap.parse_args()

    raw = sys.stdin.read() if a.input == "-" else Path(a.input).read_text(encoding="utf-8")
    text, spec = raw, ""
    try:
        data = json.loads(raw)
        t, sp = pick_outputs(data)
        if t:
            text, spec = t, sp
    except Exception:
        pass
    if a.spec:
        spec = Path(a.spec).read_text(encoding="utf-8")

    out = build_docx(text, spec, a.out, keep_markers=not a.strip_markers)
    print("已產生：", out)
    for flag, fmt in ((a.also_doc, "doc"), (a.also_pdf, "pdf")):
        if flag:
            p, err = _soffice_convert(out, fmt)
            print("已產生：" + p if p else err)


if __name__ == "__main__":
    main()
