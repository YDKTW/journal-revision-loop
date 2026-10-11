#!/usr/bin/env python3
"""
期刊稿件審查與修訂：本機一鍵執行與 Word 匯出

用途
  1. 上傳稿件（與選用的期刊格式檔）到你的 Dify 工作流並執行；
  2. 取回結果，產生 Word（.docx）報告：最終稿件、最後一輪修訂以紅字標示、待作者事項以黃底標示，
     並附逐輪審查紀錄、未解決意見、修訂對照表與系統提示。
  Word 檔可直接上傳 Google 雲端硬碟，以「Google 文件」開啟。

事前準備
  pip install requests python-docx
  在 Dify：發佈工作流 → 左側「訪問 API」取得 API 伺服器網址與 API 金鑰，設為環境變數：
    export DIFY_BASE_URL="https://你的dify網址/v1"
    export DIFY_API_KEY="app-xxxxxxxx"

用法
  python run_review.py 稿件.docx --type teaching_practice --max-rounds 2 --pass-rounds 1
  python run_review.py 稿件.pdf --type historical --format-file 某期刊稿約.pdf
  python run_review.py 稿件.md  --type content_analysis            # .md/.txt 會直接以文字送出
  python run_review.py --from-json outputs.json                    # 不呼叫 Dify，只把已存的結果轉成 Word

論文類型：quantitative, qualitative, theoretical, teaching_practice, mixed,
          content_analysis, historical, institutional_comparison,
          cross_cultural, literature_linguistics
"""
import argparse
import json
import os
import sys
from pathlib import Path

TEXT_EXT = {".txt", ".md"}
FILE_EXT = {".pdf", ".docx", ".doc"}


# ------------------------------------------------------------------ Dify API
def _headers():
    key = os.environ.get("DIFY_API_KEY")
    if not key:
        raise SystemExit("請先設定 DIFY_API_KEY（Dify 的「訪問 API」頁面可取得）。")
    return {"Authorization": "Bearer " + key}


def _base():
    b = os.environ.get("DIFY_BASE_URL", "").rstrip("/")
    if not b:
        sys.exit("請先設定環境變數 DIFY_BASE_URL，例如 https://你的dify網址/v1")
    return b


_MIME = {".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         ".doc": "application/msword", ".txt": "text/plain", ".md": "text/markdown"}


def upload(path, user):
    """上傳檔案到 Dify。檔名一律改為純英數（保留副檔名），避免中文／括號／全形字元
    造成 Dify 判斷不出副檔名而回 415 unsupported_file_type。"""
    import requests
    ext = Path(path).suffix.lower()
    if ext == ".doc":
        raise RuntimeError("Dify 不接受舊版 .doc，請在 Word 另存為 .docx（或 PDF）後再上傳。")
    safe = "upload" + ext
    with open(path, "rb") as f:
        r = requests.post(_base() + "/files/upload", headers=_headers(), data={"user": user},
                          files={"file": (safe, f, _MIME.get(ext, "application/octet-stream"))}, timeout=120)
    if r.status_code >= 300:
        raise RuntimeError("檔案上傳失敗（%s）：%s（副檔名：%s）" % (r.status_code, r.text[:300], ext or "無"))
    return r.json()["id"]


def run_workflow(inputs, user, log=print):
    """以串流模式執行（避免長時間等待逾時），回傳 (status, outputs, error)。"""
    import requests
    body = {"inputs": inputs, "response_mode": "streaming", "user": user}
    h = dict(_headers(), **{"Content-Type": "application/json"})
    with requests.post(_base() + "/workflows/run", headers=h, json=body, stream=True, timeout=(15, 900)) as r:
        if r.status_code >= 300:
            sys.exit("執行失敗（%s）：%s" % (r.status_code, r.text[:400]))
        # 以位元組逐行讀取再自行以 UTF-8 解碼：若伺服器沒有標示 charset，requests 會當成 ISO-8859-1，
        # 造成中文亂碼，而且 str.splitlines() 會在 \x85 等字元處把一行 JSON 切斷，導致完成事件遺失。
        for raw_b in r.iter_lines():
            raw = raw_b.decode("utf-8", "replace")
            if not raw or not raw.startswith("data:"):
                continue
            try:
                ev = json.loads(raw[5:].strip())
            except Exception:
                continue
            kind = ev.get("event")
            data = ev.get("data") or {}
            if kind == "node_started":
                log("  ▶ %s" % data.get("title", ""))
            elif kind == "node_finished" and data.get("status") == "failed":
                log("  ✖ 節點失敗：%s｜%s" % (data.get("title", ""), (data.get("error") or "")[:400]))
            elif kind == "workflow_finished":
                return data.get("status"), data.get("outputs") or {}, data.get("error")
            elif kind == "error":
                return "failed", {}, ev.get("message")
    return "failed", {}, "串流結束但沒有收到 workflow_finished 事件"


# ------------------------------------------------------------------ Word 匯出
def _export_manuscript(outputs, out_path):
    """依期刊格式（或內建預設）輸出修訂稿 Word，並盡量另存 .doc 與 PDF（需 LibreOffice）。"""
    try:
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        import export_docx as ed
        text = outputs.get("final_manuscript", "") if isinstance(outputs, dict) else ""
        if not text:
            print("沒有 final_manuscript，略過修訂稿輸出。")
            return
        ed.build_docx(text, outputs.get("layout_spec", ""), out_path)
        print("已產生修訂稿 Word（依期刊格式）：", out_path)
        for fmt in ("doc", "pdf"):
            p, err = ed._soffice_convert(out_path, fmt)
            if p:
                print("已產生：", p)
    except Exception as e:  # noqa: BLE001
        print("修訂稿輸出失敗（不影響報告）：", e)


def export_docx(report, out_path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")

    def east_asia(r):
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "標楷體")

    def run(par, text, color=None, hl=None, bold=False, size=None):
        r = par.add_run(text)
        r.bold = bold
        r.font.name = "Times New Roman"
        east_asia(r)
        if size:
            r.font.size = Pt(size)
        if color == "red":
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        if hl == "todo":
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return r

    def para(text="", bold=False, size=None, style=None):
        p = doc.add_paragraph(style=style)
        if text:
            run(p, text, bold=bold, size=size)
        return p

    meta = report.get("meta", {})
    h = doc.add_heading("審查與修訂報告", level=0)
    for r in h.runs:
        east_asia(r)
    para("論文類型：%s　｜　規範來源：%s　｜　結束原因：%s（%s）" % (
        meta.get("paper_type", ""), meta.get("journal_source", ""), meta.get("decision", ""), meta.get("reason", "")), size=10)
    p = para()
    run(p, meta.get("change_summary", ""), size=10)
    p = para()
    run(p, "圖例：", size=10, bold=True)
    run(p, "紅字", color="red", size=10)
    run(p, "＝最後一輪修訂中修改或新增的文字；", size=10)
    run(p, "黃底", hl="todo", size=10)
    run(p, "＝待作者補充或確認的事項。", size=10)

    fo = (report.get("final_opinion") or "").strip()
    if fo:
        for line in fo.split("\n"):
            ln = line.rstrip()
            if not ln.strip():
                continue
            if ln.startswith("## "):
                doc.add_heading(ln[3:].strip(), level=1)
            elif ln.startswith("### "):
                doc.add_heading(ln[4:].strip(), level=2)
            elif ln.lstrip().startswith(("- ", "* ")):
                p = doc.add_paragraph(style="List Bullet")
                run(p, ln.lstrip()[2:].replace("**", ""), size=11)
            else:
                para(ln.replace("**", ""), size=11)

    doc.add_heading("稿件（最終版）", level=1)
    for b in report.get("manuscript", []):
        k = b.get("k")
        if k == "table":
            rows = b.get("rows", [])
            if not rows:
                continue
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j in range(ncol):
                    cell = t.cell(i, j)
                    cell.text = ""
                    run(cell.paragraphs[0], r[j] if j < len(r) else "", bold=(i == 0), size=11)
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        lvl = b.get("l", 2) if k == "h" else 0
        size = {1: 16, 2: 14, 3: 12}.get(lvl)
        if lvl == 1:
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if "runs" in b:
            for rr in b["runs"]:
                run(p, rr["t"], color="red" if rr["c"] == "red" else None,
                    hl="todo" if rr["c"] == "todo" else None, bold=(k == "h"), size=size)
        else:
            run(p, b.get("t", ""), bold=(k == "h"), size=size)

    doc.add_page_break()
    doc.add_heading("附錄A　逐輪審查紀錄", level=1)
    for line in report.get("history", []):
        doc.add_paragraph(line, style="List Bullet")
    doc.add_heading("附錄B　最後仍未解決的意見", level=1)
    un = report.get("unresolved", [])
    if not un:
        para("（無）")
    for u in un:
        p = doc.add_paragraph(style="List Bullet")
        run(p, "[%s｜%s｜%s] " % (u.get("id", ""), u.get("severity", ""), u.get("location", "")), bold=True, size=11)
        run(p, u.get("problem", ""), size=11)
        if u.get("suggestion"):
            run(p, "　建議：" + u["suggestion"], size=11)
    doc.add_heading("附錄C　最後一輪修訂對照表", level=1)
    for line in (report.get("response_table") or "（無）").split("\n"):
        para(line, size=11)
    doc.add_heading("附錄D　系統提示", level=1)
    for title, key in (("前處理報告", "preprocess_report"), ("格式整理報告", "format_report"), ("警示與提示", "flags")):
        if report.get(key):
            para(title, bold=True, size=11)
            for line in report[key].split("\n"):
                para(line, size=10)
    doc.save(out_path)


# ------------------------------------------------------------------ 主程式
def main():
    ap = argparse.ArgumentParser(description="期刊稿件審查與修訂（Dify）＋ Word 匯出")
    ap.add_argument("manuscript", nargs="?", help="稿件檔（.docx .pdf .md .txt）")
    ap.add_argument("--type", default="teaching_practice", help="論文類型（見檔案開頭說明）")
    ap.add_argument("--max-rounds", type=int, default=4)
    ap.add_argument("--pass-rounds", type=int, default=2)
    ap.add_argument("--format-file", help="期刊格式檔（稿約）；未提供則使用內建預設格式")
    ap.add_argument("--known-gaps", default="", help="作者已知待補事項（選填；多項請用分號分隔）。列在此處的缺口，審查者的嚴重度最高為 minor")
    ap.add_argument("--out", help="輸出 Word 檔路徑（預設：與稿件同名加 _審查報告.docx）")
    ap.add_argument("--from-json", help="不呼叫 Dify，直接用已存的結果 JSON 產生 Word")
    ap.add_argument("--user", default="journal-review-cli")
    a = ap.parse_args()

    if a.from_json:
        data = json.loads(Path(a.from_json).read_text(encoding="utf-8"))
        rep = data.get("report_json", data)
        report = json.loads(rep) if isinstance(rep, str) else rep
        out = a.out or str(Path(a.from_json).with_suffix(".docx"))
        export_docx(report, out)
        print("已產生：", out)
        _export_manuscript(data, str(Path(out).with_name(Path(out).stem + "_修訂稿_依期刊格式.docx")))
        return

    if not a.manuscript:
        ap.error("請提供稿件檔，或使用 --from-json")
    mp = Path(a.manuscript)
    if not mp.exists():
        sys.exit("找不到稿件檔：%s" % mp)

    inputs = {"paper_type": a.type, "max_rounds": a.max_rounds, "required_pass_rounds": a.pass_rounds,
              "format_source": "內建預設格式", "manuscript": "", "known_gaps": a.known_gaps.replace(";", "\n").replace("；", "\n")}
    if mp.suffix.lower() in TEXT_EXT:
        inputs["manuscript"] = mp.read_text(encoding="utf-8")
    elif mp.suffix.lower() in FILE_EXT:
        print("上傳稿件…")
        inputs["file"] = {"type": "document", "transfer_method": "local_file", "upload_file_id": upload(mp, a.user)}
    else:
        sys.exit("不支援的稿件格式：%s" % mp.suffix)
    if a.format_file:
        print("上傳期刊格式檔…")
        inputs["format_source"] = "上傳期刊格式檔"
        inputs["journal_file"] = {"type": "document", "transfer_method": "local_file",
                                  "upload_file_id": upload(a.format_file, a.user)}

    print("執行工作流（依稿件長度，可能需要數分鐘）…")
    status, outputs, err = run_workflow(inputs, a.user)
    raw_path = mp.with_name(mp.stem + "_dify輸出.json")
    raw_path.write_text(json.dumps(outputs, ensure_ascii=False, indent=2), encoding="utf-8")
    print("原始輸出已存：", raw_path)
    if status != "succeeded":
        sys.exit("工作流未成功（%s）：%s" % (status, err))

    report = json.loads(outputs["report_json"])
    out = a.out or str(mp.with_name(mp.stem + "_審查報告.docx"))
    export_docx(report, out)
    print("\n結束原因：%s（%s）" % (outputs.get("decision"), outputs.get("last_reason")))
    print(outputs.get("change_summary", ""))
    print("已產生 Word 報告：", out)
    _export_manuscript(outputs, str(mp.with_name(mp.stem + "_修訂稿_依期刊格式.docx")))


if __name__ == "__main__":
    main()
